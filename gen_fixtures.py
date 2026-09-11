"""Генератор тестовых файлов для smoke-теста (test_data)."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_data")
os.makedirs(OUT, exist_ok=True)

# 1. TXT
with open(os.path.join(OUT, "проект_ирис.txt"), "w", encoding="utf-8") as f:
    f.write("Проект ИРИС (2024 год).\n"
            "В проекте ИРИС для Газпром нефтехим Салават использовался 1С:Документооборот "
            "для согласования договоров подряда. Архив согласования хранится в базе "
            "docob-iris. Дополнительно развернут модуль интеграции с 1С:ERP.\n")

# 2. CSV
with open(os.path.join(OUT, "реестр_систем.csv"), "w", encoding="utf-8") as f:
    f.write("проект;система;год;роль\n")
    f.write("ИРИС;1С:Документооборот;2024;согласование договоров\n")
    f.write("РТИ;1С:Документооборот;2023;архив входящих писем\n")
    f.write("Умный склад;1С:WMS;2025;логистика\n")

# 3. PDF
import fitz

doc = fitz.open()
page = doc.new_page()
page.insert_text((72, 72), "Отчет по проекту Цифровой склад, 2025",
                 fontsize=14, fontname="helv")
page.insert_text((72, 100),
                 "Использовался 1С:Документооборот версии 2.1 для регламента СОУД и "
                 "маршрутов согласования проектной документации.")
doc.save(os.path.join(OUT, "отчет_цс.pdf"))
doc.close()

# 4. DOCX
from docx import Document

d = Document()
d.add_heading("Пояснительная записка проекта РТК", 1)
d.add_paragraph(
    "На этапе предпроектного обследования было принято решение использовать "
    "1С:Документооборот совместно с MS Project для управления календарными планами. "
    "Ответственный: Саша. Сроки: 2024-2026.")
d.save(os.path.join(OUT, "записка_ртк.docx"))

# 5. XLSX
import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Системы"
ws.append(["Проект", "Система", "Статус"])
ws.append(["ИРИС", "1С:Документооборот", "внедрен"])
ws.append(["РТК", "1С:Документооборот", "внедрен"])
ws.append(["ЦС", "Confluence", "внедрена"])
wb.save(os.path.join(OUT, "реестр_систем.xlsx"))

# 6. PPTX
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Презентация проекта ЦС"
slide.placeholders[1].text = (
    "В проекте ЦС (Цифровой склад) использовался 1С:Документооборот для "
    "согласования КД и документации ЕСКД.")
prs.save(os.path.join(OUT, "презентация_цс.pptx"))

print("Созданы тестовые файлы в %s" % OUT)
for fn in sorted(os.listdir(OUT)):
    print(" -", fn)