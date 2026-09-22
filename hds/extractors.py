"""Извлечение текста: txt/md/код, PDF, DOCX, XLSX, PPTX (OCR пустых страниц PDF).

md и DOCX возвращают сегменты с полем head — «путь заголовков» секции
(например «# Раздел / ## Подраздел»); структурный чанкер (hds/chunker.py)
добавляет его в начало каждого чанка секции."""
import os
import re
import sys

from .config import dig

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".log", ".csv", ".tsv", ".json", ".jsonl", ".xml",
    ".yaml", ".yml", ".ini", ".cfg", ".conf", ".env", ".bat", ".cmd", ".ps1",
    ".py", ".pyw", ".js", ".mjs", ".ts", ".java", ".cs", ".cpp", ".c", ".h", ".hpp",
    ".sql", ".html", ".htm", ".css", ".php", ".rb", ".go", ".rs", ".sh", ".vbs",
    ".reg", ".url", ".srt", ".ass", ".tex", ".rst", ".rtf",
}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx", ".xlsm"}
PPTX_EXTS = {".pptx"}

_OCR_OK = None


def kind_for_ext(ext):
    ext = (ext or "").lower()
    if ext in PDF_EXTS:
        return "pdf"
    if ext in DOCX_EXTS:
        return "docx"
    if ext in XLSX_EXTS:
        return "xlsx"
    if ext in PPTX_EXTS:
        return "pptx"
    if ext in TEXT_EXTS:
        return "text"
    return None


def seg(text, page=None, t_start=None, t_end=None, head=None):
    d = {"text": text, "page": page, "t_start": t_start, "t_end": t_end}
    if head:
        d["head"] = head  # путь заголовков секции — для структурного чанкера
    return d


def read_text_file(path):
    with open(path, "rb") as f:
        data = f.read()
    text = None
    for enc in ("utf-8", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode("utf-8", "ignore")
    return text.replace("\x00", "")


def _tesseract_ready(cfg):
    ok = _tesseract_ready.__dict__.get("ok")
    if ok is None:
        try:
            import pytesseract

            cmd = dig(cfg, "index.ocr_tesseract_cmd", "") or None
            if cmd:
                pytesseract.pytesseract.tesseract_cmd = cmd
            # языковые пакеты, установленные без прав администратора, лежат в
            # пользовательском tessdata — указываем его, если он существует
            user_td = os.path.join(os.environ.get("LOCALAPPDATA", ""),
                                   "Tesseract-OCR", "tessdata")
            if sys.platform == "darwin" and not os.path.isdir(user_td):
                # macOS: tesseract из Homebrew — языки в share/tessdata
                for td in ("/opt/homebrew/share/tessdata",
                           "/usr/local/share/tessdata"):
                    if os.path.isdir(td):
                        user_td = td
                        break
            if os.path.isdir(user_td):
                os.environ.setdefault("TESSDATA_PREFIX", user_td)
            pytesseract.get_tesseract_version()
            ok = True
        except Exception:  # noqa: BLE001
            ok = False
        _tesseract_ready.__dict__["ok"] = ok
    return ok


def _ocr_pil_image(img, cfg):
    import pytesseract

    return pytesseract.image_to_string(
        img, lang=dig(cfg, "index.ocr_lang", "rus+eng")
    )


_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_DOCX_HEADING_RE = re.compile(r"^(?:heading|заголовок)\s+(\d+)", re.IGNORECASE)


def _head_path(head_parts):
    """Список [(level, title)] -> строка «# Раздел / ## Подраздел»."""
    return " / ".join("%s %s" % ("#" * lvl, title) for lvl, title in head_parts)


def _push_heading(head_parts, level, title):
    """Заголовок уровня level: убрать более глубокие/равные, добавить новый."""
    while head_parts and head_parts[-1][0] >= level:
        head_parts.pop()
    head_parts.append((level, title))


def extract_markdown(path, cfg):
    """Markdown: сегмент на секцию (между заголовками), head — путь заголовков."""
    text = read_text_file(path)
    segs, cur = [], []
    head_parts = []

    def flush():
        body = "\n".join(cur).strip()
        if body:
            segs.append(seg(body, head=_head_path(head_parts) if head_parts else None))
        del cur[:]

    for line in text.splitlines():
        m = _MD_HEADING_RE.match(line)
        if m:
            flush()
            _push_heading(head_parts, len(m.group(1)), m.group(2).strip())
            continue
        cur.append(line)
    flush()
    return segs if segs else [seg(text)]


def extract_pdf(path, cfg):
    import fitz

    use_ocr = dig(cfg, "index.ocr", True) and _tesseract_ready(cfg)
    doc = fitz.open(path)
    segs = []
    try:
        for i, page in enumerate(doc, start=1):
            txt = page.get_text().strip()
            if len(txt) < 40 and use_ocr:
                try:
                    from PIL import Image

                    pix = page.get_pixmap(dpi=200)
                    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    txt = _ocr_pil_image(img, cfg).strip() or txt
                except Exception:  # noqa: BLE001
                    pass
            if txt:
                segs.append(seg(txt, page=i))
    finally:
        doc.close()
    return segs


def extract_docx(path, cfg):
    """DOCX: сегмент на секцию заголовков (head — путь заголовков) + таблицы."""
    from docx import Document

    d = Document(path)
    segs, cur = [], []
    head_parts = []

    def flush():
        body = "\n".join(cur).strip()
        if body:
            segs.append(seg(body, head=_head_path(head_parts) if head_parts else None))
        del cur[:]

    for p in d.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        m = _DOCX_HEADING_RE.match(style.strip())
        if m:
            flush()
            try:
                level = min(6, max(1, int(m.group(1))))
            except ValueError:
                level = 1
            _push_heading(head_parts, level, t)
            continue
        if style.strip().lower() == "title":
            flush()
            _push_heading(head_parts, 1, t)
            continue
        cur.append(p.text)
    flush()
    for tbl in d.tables:
        rows = []
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            segs.append(seg("\n".join(rows),
                            head=_head_path(head_parts) if head_parts else None))
    return segs if segs else [seg("")]


def extract_xlsx(path, cfg):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    segs = []
    try:
        for idx, ws in enumerate(wb.worksheets, start=1):
            lines = []
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                if r >= 5000:
                    lines.append("... (лист обрезан)")
                    break
                vals = [str(v) for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append(" | ".join(vals))
            if lines:
                segs.append(seg("Лист: %s\n%s" % (ws.title, "\n".join(lines)), page=idx))
    finally:
        wb.close()
    return segs


def extract_pptx(path, cfg):
    from pptx import Presentation

    prs = Presentation(path)
    segs = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            segs.append(seg("\n".join(parts), page=i))
    return segs


def extract_text(path, cfg):
    return [seg(read_text_file(path))]


EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "xlsx": extract_xlsx,
    "pptx": extract_pptx,
    "text": extract_text,
}


def extract(path, cfg, progress_cb=None):
    """Возвращает (kind, segments). Для неизвестных форматов -> (None, []).
    progress_cb вызывается с процентом обработки медиафайлов (0–100)."""
    from .extract_av import kind_for_ext_media as _av_kind
    from .extract_av import extract_dispatch_media
    from .extract_static import kind_for_ext_media as _static_kind
    from .extract_static import extract_dispatch_static

    ext = os.path.splitext(path)[1].lower()
    kind = kind_for_ext(ext) or _static_kind(ext) or _av_kind(ext)
    if kind is None:
        return None, []
    if kind == "text" and ext in (".md", ".markdown"):
        return kind, extract_markdown(path, cfg)
    if kind in EXTRACTORS:
        return kind, EXTRACTORS[kind](path, cfg)
    fn = extract_dispatch_static if kind in ("mpp", "image") else extract_dispatch_media
    if kind == "media":
        return fn(path, cfg, kind, progress_cb)
    return fn(path, cfg, kind)