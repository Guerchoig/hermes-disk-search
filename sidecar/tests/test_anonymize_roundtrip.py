"""Round-trip по форматам (P1): TXT/MD — байт-в-байт после «анон → деан»;
DOCX/XLSX — текст по сегментам + структура + ANON_MARKER; частичные
сценарии; FILE_BUSY; дедуп; отсутствие PII в логах.

NER-модели в тестах НЕ нужны: сценарии идут через детерминированный
`anonymize_values` (значения задаются явно). Полный `anonymize` без весов
проверяется в test_anonymize_worker (kind: model_missing).

Запуск: python -m unittest discover -s sidecar/tests -t sidecar -v
"""
from __future__ import annotations

import asyncio
import io
import logging
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from hds_anonymize.file_parser import (  # noqa: E402
    ANON_MARKER,
    FileParser,
    read_anon_marker,
)
from hds_anonymize.handlers import (  # noqa: E402
    AnonymizeHandlers,
    FileBusyError,
)
from hds_anonymize.mapping_store import MappingStore  # noqa: E402

VALUES = ["Иванов Иван", "+7 912 345-67-89", "51 903,14"]

TXT_ORIGINAL = (
    "Договор №0095/23\r\n"
    "Исполнитель: Иванов Иван (тел. +7 912 345-67-89)\r\n"
    "Сумма: 51 903,14 руб.\r\n"
    "Сумма: 51 903,14 руб. — повтор\r\n"
)


def _run(coro):
    return asyncio.run(coro)


def _make_docx(path: str) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Исполнитель: Иванов Иван")
    d.add_paragraph("Телефон: +7 912 345-67-89")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Контрагент"
    t.cell(0, 1).text = "ООО Ромашка"
    t.cell(1, 0).text = "Сумма"
    t.cell(1, 1).text = "51 903,14"
    d.save(path)


def _make_xlsx(path: str) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws["A1"] = "Контрагент"
    ws["B1"] = "ООО Ромашка"
    ws["A2"] = "Контакт"
    ws["B2"] = "Иванов Иван"
    ws["A3"] = "Сумма"
    ws["B3"] = 386887.42  # числовая ячейка: структура сохраняется
    ws["A4"] = "Телефон"
    ws["B4"] = "+7 912 345-67-89"
    wb.save(path)


class RoundtripBase(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory(prefix="hds-anon-rt-")
        self.dir = self.tmp.name
        self.db = os.path.join(self.dir, "anonymize.db")
        self.store = MappingStore(self.db)
        self.h = AnonymizeHandlers(self.store)

    def tearDown(self) -> None:
        self.store.close()
        self.tmp.cleanup()

    def anon(self, src: str, values=None):
        out = os.path.join(
            self.dir,
            os.path.basename(src).rsplit(".", 1)[0] + ".anonymized."
            + src.rsplit(".", 1)[1],
        )
        res = _run(self.h.anonymize(src, out, values=values or VALUES))
        return out, res

    def deanon(self, src: str, tokens=None):
        stem = os.path.basename(src).rsplit(".", 1)[0]
        if stem.endswith(".anonymized"):
            stem = stem[: -len(".anonymized")]
        ext = src.rsplit(".", 1)[1]
        out = os.path.join(self.dir, f"{stem}.deanon.{ext}")
        res = _run(self.h.deanonymize(src, out, tokens=tokens))
        return out, res

class TextRoundtripTest(RoundtripBase):
    """TXT/MD: байт-в-байт после «анон → деан» (приёмка §5)."""

    def test_txt_byte_exact_roundtrip(self):
        src = os.path.join(self.dir, "doc.txt")
        with open(src, "wb") as f:
            f.write(TXT_ORIGINAL.encode("utf-8"))
        anon_out, res = self.anon(src)
        anon_bytes = open(anon_out, "rb").read()
        # значения ушли, токены появились; построчная структура сохранена
        self.assertNotIn("Иванов Иван".encode(), anon_bytes)
        self.assertNotIn("+7 912 345-67-89".encode(), anon_bytes)
        self.assertIn("[PERSON_1]", anon_bytes.decode("utf-8"))
        self.assertIn("[MISC_1]", anon_bytes.decode("utf-8"))
        # дедуп: значение дважды — один токен (count=2)
        money = [t for t in res["tokens"] if t["token"] == "[MISC_2]"]
        self.assertEqual(money, [{"token": "[MISC_2]", "type": "MISC", "count": 2}])

        deanon_out, _res = self.deanon(anon_out)
        self.assertEqual(open(deanon_out, "rb").read(), TXT_ORIGINAL.encode("utf-8"))

    def test_md_byte_exact_roundtrip(self):
        src = os.path.join(self.dir, "doc.md")
        original = "# Заголовок\n\nКонтакт: Иванов Иван, 51 903,14\n".encode("utf-8")
        with open(src, "wb") as f:
            f.write(original)
        anon_out, _ = self.anon(src)
        deanon_out, _ = self.deanon(anon_out)
        self.assertEqual(open(deanon_out, "rb").read(), original)


class OfficeRoundtripTest(RoundtripBase):
    """DOCX/XLSX: маркер, отсутствие PII, текст и структура после деанона."""

    def test_docx_roundtrip(self):
        src = os.path.join(self.dir, "doc.docx")
        _make_docx(src)
        parsed0 = asyncio.run(FileParser().parse(open(src, "rb").read(), "doc.docx"))
        self.assertIn("Иванов Иван", parsed0.text)

        anon_out, _ = self.anon(src, values=["Иванов Иван", "ООО Ромашка"])
        # маркер анонимизации в DOCX (docProps/comments)
        self.assertEqual(read_anon_marker(anon_out), ANON_MARKER)
        anon_parsed = asyncio.run(
            FileParser().parse(open(anon_out, "rb").read(), "doc.docx")
        )
        self.assertNotIn("Иванов Иван", anon_parsed.text)
        self.assertNotIn("ООО Ромашка", anon_parsed.text)
        self.assertIn("[PERSON_1]", anon_parsed.text)

        deanon_out, res = self.deanon(anon_out)
        self.assertEqual(res["unmatched"], [])
        de_parsed = asyncio.run(
            FileParser().parse(open(deanon_out, "rb").read(), "doc.docx")
        )
        self.assertEqual(de_parsed.text, parsed0.text)
        self.assertEqual(de_parsed.structure, parsed0.structure)

    def test_xlsx_roundtrip_structure_and_marker(self):
        src = os.path.join(self.dir, "doc.xlsx")
        _make_xlsx(src)
        parsed0 = asyncio.run(FileParser().parse(open(src, "rb").read(), "doc.xlsx"))

        anon_out, _ = self.anon(
            src, values=["Иванов Иван", "ООО Ромашка", "+7 912 345-67-89"]
        )
        self.assertEqual(read_anon_marker(anon_out), ANON_MARKER)
        anon_parsed = asyncio.run(
            FileParser().parse(open(anon_out, "rb").read(), "doc.xlsx")
        )
        self.assertNotIn("Иванов Иван", anon_parsed.text)
        self.assertNotIn("ООО Ромашка", anon_parsed.text)

        deanon_out, _ = self.deanon(anon_out)
        de_parsed = asyncio.run(
            FileParser().parse(open(deanon_out, "rb").read(), "doc.xlsx")
        )
        self.assertEqual(de_parsed.text, parsed0.text)
        self.assertEqual(de_parsed.structure, parsed0.structure)

        # числовая ячейка осталась числом (не строкой)
        import openpyxl

        wb = openpyxl.load_workbook(deanon_out)
        self.assertIsInstance(wb["Лист1"]["B3"].value, (int, float))
        self.assertAlmostEqual(wb["Лист1"]["B3"].value, 386887.42)

class PartialScenariosTest(RoundtripBase):
    """Частичные сценарии: выборочная анонимизация/деанонимизация, busy,
    сквозная нумерация, inspect, логи."""

    def test_values_unmatched_reported(self):
        src = os.path.join(self.dir, "doc.txt")
        with open(src, "w", encoding="utf-8", newline="") as f:
            f.write("Контакт: Иванов Иван\n")
        _, res = self.anon(src, values=["Иванов Иван", "Нет Такого"])
        self.assertIn("Иванов Иван", res["matched"])
        self.assertEqual(res["unmatched"], ["Нет Такого"])

    def test_deanonymize_values_subset(self):
        src = os.path.join(self.dir, "doc.txt")
        with open(src, "wb") as f:
            f.write(TXT_ORIGINAL.encode("utf-8"))
        anon_out, res = self.anon(src)
        tokens = [t["token"] for t in res["tokens"]]
        subset = tokens[:1]
        deanon_out, dres = self.deanon(anon_out, tokens=subset)
        text = open(deanon_out, "r", encoding="utf-8").read()
        # восстановлено только выбранное; остальные токены на месте
        for t in tokens:
            if t in subset:
                self.assertNotIn(t, text)
            else:
                self.assertIn(t, text)
        # unmatched — только запрошенные токены, ОТСУТСТВУЮЩИЕ в БД;
        # все запрошенные существуют → пусто (§3)
        self.assertEqual(dres["unmatched"], [])
        self.assertEqual(sorted(dres["restored"]), sorted(subset))

    def test_deanonymize_values_unmatched(self):
        src = os.path.join(self.dir, "doc.txt")
        with open(src, "wb") as f:
            f.write(TXT_ORIGINAL.encode("utf-8"))
        anon_out, _ = self.anon(src)
        _, dres = self.deanon(anon_out, tokens=["[PERSON_777]"])
        self.assertEqual(dres["unmatched"], ["[PERSON_777]"])

    def test_inspect_without_values(self):
        src = os.path.join(self.dir, "doc.txt")
        with open(src, "wb") as f:
            f.write(TXT_ORIGINAL.encode("utf-8"))
        anon_out, res = self.anon(src)
        info = _run(self.h.inspect(anon_out))
        self.assertEqual(info["marker_version"], None)  # TXT без маркера
        tokens = {t["token"]: t["count"] for t in res["tokens"]}
        got = {p["token"]: p["count"] for p in info["placeholders"]}
        self.assertEqual(got, tokens)
        self.assertTrue(all(r["in_db"] for r in info["resolved"]))
        # PII-значений в inspect нет
        self.assertNotIn("Иванов", str(info))
        self.assertNotIn("+7 912", str(info))

    def test_numbering_is_continuous_across_files(self):
        """Сквозная нумерация по актам: второй файл продолжает счётчик."""
        src1 = os.path.join(self.dir, "one.txt")
        src2 = os.path.join(self.dir, "two.txt")
        with open(src1, "w", encoding="utf-8", newline="") as f:
            f.write("Сотрудник: Петров Пётр\n")
        with open(src2, "w", encoding="utf-8", newline="") as f:
            f.write("Сотрудник: Сидоров Сидор\n")
        _, r1 = self.anon(src1, values=["Петров Пётр"])
        _, r2 = self.anon(src2, values=["Сидоров Сидор"])
        self.assertEqual([t["token"] for t in r1["tokens"]], ["[PERSON_1]"])
        self.assertEqual([t["token"] for t in r2["tokens"]], ["[PERSON_2]"])

    def test_out_dir_created_on_first_act(self):
        """Выходной каталог создаётся при первом акте (data/anonymized)."""
        src = os.path.join(self.dir, "doc.txt")
        with open(src, "w", encoding="utf-8", newline="") as f:
            f.write("Контакт: Иванов Иван\n")
        out_dir = os.path.join(self.dir, "nested", "out")
        out = os.path.join(out_dir, "doc.anonymized.txt")
        res = _run(self.h.anonymize(src, out, values=["Иванов Иван"]))
        self.assertTrue(os.path.isfile(res["out_path"]))

    @unittest.skipUnless(
        os.name == "nt",
        "FILE_BUSY — сценарий Windows: Word/LibreOffice держат файл и замена "
        "даёт PermissionError. На POSIX rename поверх открытого файла разрешён, "
        "поэтому сценарий там не воспроизводится",
    )
    def test_file_busy(self):
        src = os.path.join(self.dir, "doc.txt")
        with open(src, "w", encoding="utf-8", newline="") as f:
            f.write("Контакт: Иванов Иван\n")
        out = os.path.join(self.dir, "busy.anonymized.txt")
        with open(out, "wb") as locked:
            locked.write(b"")
            locked.flush()
            with self.assertRaises(FileBusyError):
                _run(self.h.anonymize(src, out, values=["Иванов Иван"]))

    def test_already_anonymized_warning(self):
        src = os.path.join(self.dir, "doc.txt")
        with open(src, "w", encoding="utf-8", newline="") as f:
            f.write("Контакт: Иванов Иван\n")
        out1 = os.path.join(self.dir, "a.anonymized.txt")
        _run(self.h.anonymize(src, out1, values=["Иванов Иван"]))
        res2 = _run(self.h.anonymize(
            out1, os.path.join(self.dir, "b.anonymized.txt"),
            values=["Иванов Иван"]))
        self.assertTrue(res2["warnings"], "повторная анонимизация даёт warning")

    def test_no_pii_in_logs(self):
        """В логи не попадают значения (только токены/счётчики)."""
        captured = io.StringIO()
        handler = logging.StreamHandler(captured)
        lg = logging.getLogger("hds_anonymize")
        lg.addHandler(handler)
        try:
            src = os.path.join(self.dir, "doc.txt")
            with open(src, "w", encoding="utf-8", newline="") as f:
                f.write("Контакт: Иванов Иван, +7 912 345-67-89\n")
            self.anon(src)
        finally:
            lg.removeHandler(handler)
        logs = captured.getvalue()
        for value in VALUES:
            self.assertNotIn(value, logs, f"PII утекло в лог: {value!r}")


if __name__ == "__main__":
    unittest.main()
