"""Тесты извлечения текста из форматов: txt, csv, pdf, docx, xlsx, pptx, mpp, exe."""
import os
import sys
import tempfile
import unittest

from helpers import write_text  # noqa: I100

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from hds.config import load  # noqa: E402
from hds.extractors import extract  # noqa: E402


class ExtractorTests(unittest.TestCase):
    def setUp(self):
        self.cfg = load()
        self.tmp = tempfile.mkdtemp(prefix="hds-ext-")

    def test_text(self):
        p = write_text(self.tmp, "a.txt", "строка про 1С:Документооборот")
        kind, segs = extract(p, self.cfg)
        self.assertEqual(kind, "text")
        self.assertIn("Документооборот", segs[0]["text"])

    def test_csv(self):
        p = write_text(self.tmp, "a.csv", "проект;система\nИРИС;1С:Документооборот")
        kind, segs = extract(p, self.cfg)
        self.assertIn("Документооборот", segs[0]["text"])

    def test_pdf(self):
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "PDF pro document", fontname="helv")
        p = os.path.join(self.tmp, "a.pdf")
        doc.save(p)
        doc.close()
        kind, segs = extract(p, self.cfg)
        self.assertEqual(kind, "pdf")
        self.assertTrue(any("PDF" in s["text"] for s in segs))

    def test_docx(self):
        from docx import Document

        d = Document()
        d.add_paragraph("Проект РТК использовал 1С:Документооборот")
        p = os.path.join(self.tmp, "a.docx")
        d.save(p)
        kind, segs = extract(p, self.cfg)
        self.assertEqual(kind, "docx")
        self.assertIn("РТК", segs[0]["text"])

    def test_xlsx(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Проект", "Система"])
        ws.append(["ИРИС", "1С:Документооборот"])
        p = os.path.join(self.tmp, "a.xlsx")
        wb.save(p)
        kind, segs = extract(p, self.cfg)
        self.assertEqual(kind, "xlsx")
        self.assertIn("Документооборот", segs[0]["text"])

    def test_pptx(self):
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Проект ЦС"
        prs.save(os.path.join(self.tmp, "a.pptx"))
        kind, segs = extract(p := os.path.join(self.tmp, "a.pptx"), self.cfg)
        self.assertEqual(kind, "pptx")
        self.assertIn("ЦС", segs[0]["text"])

    def test_unknown_ext(self):
        p = write_text(self.tmp, "a.exe", "MZ...")
        kind, segs = extract(p, self.cfg)
        self.assertIsNone(kind)
        self.assertEqual(segs, [])

    def test_image_includes_folder(self):
        """Фича: в чанк картинки попадает имя папки (поиск фото по папкам)."""
        import tempfile
        from PIL import Image

        sub = os.path.join(self.tmp, "Цветы")
        os.makedirs(sub, exist_ok=True)
        p = os.path.join(sub, "photo.jpg")
        Image.new("RGB", (60, 40), "white").save(p)
        kind, segs = extract(p, self.cfg)
        self.assertEqual(kind, "image")
        text = segs[0]["text"]
        self.assertIn("папка: Цветы", text, "имя папки должно быть в чанке")

    def test_mpp_without_mpxj_reports(self):
        """Фича: .mpp без установленного mpxj — пометка, а не падение."""
        p = write_text(self.tmp, "a.mpp", "заглушка")
        kind, segs = extract(p, self.cfg)
        self.assertEqual(kind, "mpp")
        self.assertTrue(segs)  # сегмент с пояснением


if __name__ == "__main__":
    unittest.main()